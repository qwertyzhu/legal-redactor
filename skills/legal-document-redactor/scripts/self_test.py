#!/usr/bin/env python3
"""self_test.py — legal-document-redactor skill 安装自检 / 回归测试。

全部使用现场生成的虚构文书，不触碰任何真实文件；在临时目录运行，结束自动清理。
用法（在 skill 目录下）：
    python3 scripts/self_test.py            # 人类可读报告
    python3 scripts/self_test.py --json     # 机器可读（供 Agent 判断）

覆盖：production 默认、ai --auto-confident、指哪打哪（keep-categories 反向枚举、
点名人名）、court-style 模板替身、verify 退出码双向、Ollama 检测与识别（无 Ollama 自动 SKIP）。
退出码：全部 PASS/SKIP → 0；任何 FAIL → 1。
"""
from __future__ import annotations

import argparse
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

SKILL_DIR = Path(__file__).resolve().parent.parent
PY = sys.executable

FICTION_PARAS = [
    "原告张测三，男，1985年3月12日出生，住杭州市西湖区文三路100号，身份证号330106198503121234，手机138 0000 2222，邮箱zhang@example.com，车牌号浙A12345。",
    "被告北测文化传播有限公司，统一社会信用代码91330100MA27TEST1X，法定代表人李测四。",
    "案号：（2024）京0491民初1234号。涉案作品《星河测例》，标的金额1280000元。",
    "收款账户：6222 0012 3456 7890。",
]


def run(cmd: list[str], **kw) -> subprocess.CompletedProcess:
    return subprocess.run(cmd, capture_output=True, text=True, timeout=600, **kw)


def docx_text(path: Path) -> str:
    from docx import Document
    return "\n".join(p.text for p in Document(str(path)).paragraphs)


def make_fixture(workdir: Path) -> Path:
    from docx import Document
    d = Document()
    for p in FICTION_PARAS:
        d.add_paragraph(p)
    src = workdir / "fiction.docx"
    d.save(str(src))
    return src


class Suite:
    def __init__(self) -> None:
        self.results: list[dict] = []

    def check(self, name: str, ok: bool, detail: str = "") -> None:
        self.results.append({"name": name, "status": "PASS" if ok else "FAIL", "detail": detail})

    def skip(self, name: str, reason: str) -> None:
        self.results.append({"name": name, "status": "SKIP", "detail": reason})



def main() -> None:
    ap = argparse.ArgumentParser(description="skill 安装自检 / 回归测试（虚构数据）")
    ap.add_argument("--json", action="store_true")
    args = ap.parse_args()

    s = Suite()
    cli = shutil.which("legal-redactor")
    if not cli:
        s.check("legal-redactor 在 PATH", False, "未找到；先安装 Python 包或做软链")
        report(s, args.json)
        sys.exit(1)
    s.check("legal-redactor 在 PATH", True, cli)

    with tempfile.TemporaryDirectory(prefix="lr-selftest-") as td:
        w = Path(td)
        src = make_fixture(w)

        # T1: production 默认——证件/手机/邮箱去掉，当事人姓名与案号保留
        out = w / "t1.docx"
        r = run([cli, "redact", str(src), "--mode", "production", "--auto-confident", "-o", str(out)])
        text = docx_text(out) if out.exists() else ""
        ok = (r.returncode == 0 and "330106198503121234" not in text and "138 0000 2222" not in text
              and "zhang@example.com" not in text and "张测三" in text and "（2024）京0491民初1234号" in text)
        s.check("T1 production 默认（去高敏、留当事人和案号）", ok)

        # T2: ai --auto-confident——高置信实体去掉
        out = w / "t2.docx"
        r = run([cli, "redact", str(src), "--mode", "ai", "--auto-confident", "-o", str(out)])
        text = docx_text(out) if out.exists() else ""
        ok = r.returncode == 0 and "北测文化传播有限公司" not in text and "（2024）京0491民初1234号" not in text
        s.check("T2 ai --auto-confident（单位、案号去掉）", ok)

        # T3: 指哪打哪——只去身份证号，其余结构性字段放行
        out = w / "t3.docx"
        r = run([cli, "redact", str(src), "--mode", "production",
                 "--keep-categories", "mobile,email,bank_account,uscc,case_number,landline", "-o", str(out)])
        text = docx_text(out) if out.exists() else ""
        ok = (r.returncode == 0 and "330106198503121234" not in text
              and "138 0000 2222" in text and "zhang@example.com" in text)
        s.check("T3 指哪打哪：只脱敏身份证号", ok)

        # T4: 点名人名——只替换指定人名，其他一概不动
        ent = w / "t4.json"
        ent.write_text(json.dumps({"entities": [
            {"original": "张测三", "category": "person", "role": "other", "replacement": "某甲"}]},
            ensure_ascii=False), encoding="utf-8")
        out = w / "t4.docx"
        r = run([cli, "redact", str(src), "--mode", "production", "--entities", str(ent),
                 "--keep-categories", "id_card,mobile,landline,email,bank_account,uscc,case_number",
                 "-o", str(out)])
        text = docx_text(out) if out.exists() else ""
        ok = (r.returncode == 0 and "某甲" in text and "张测三" not in text
              and "330106198503121234" in text and "李测四" in text)
        s.check("T4 指哪打哪：只替换点名人名", ok)

        # T5: court-style 替身（保留姓氏+某、出生日期、车牌号盲区覆盖）
        ent = w / "t5.json"
        ent.write_text(json.dumps({"entities": [
            {"original": "张测三", "category": "person", "role": "party", "replacement": "张某"},
            {"original": "1985年3月12日出生", "category": "other", "replacement": "19××年×月出生"},
            {"original": "浙A12345", "category": "other", "replacement": "（车牌号已删除）"}]},
            ensure_ascii=False), encoding="utf-8")
        out = w / "t5.docx"
        r = run([cli, "redact", str(src), "--mode", "ai", "--entities", str(ent), "-o", str(out)])
        text = docx_text(out) if out.exists() else ""
        ok = (r.returncode == 0 and "张某" in text and "张测三" not in text
              and "1985年3月12日" not in text and "浙A12345" not in text)
        s.check("T5 裁判文书式替身+盲区字段", ok)

        # T6: verify 退出码双向
        ok1 = run([cli, "verify", str(w / "t5.docx"), "--mode", "ai"]).returncode == 0
        ok2 = run([cli, "verify", str(src), "--mode", "ai"]).returncode != 0
        s.check("T6 verify 退出码（脱敏品过 / 原文不过）", ok1 and ok2)

        # T7: ollama_setup 可运行、JSON 可解析
        r = run([PY, str(SKILL_DIR / "scripts" / "ollama_setup.py"), "--json"])
        try:
            info = json.loads(r.stdout)
            s.check("T7 ollama_setup --json", r.returncode == 0,
                    f"ollama={info['ollama_installed']} ram={info['ram_gb']}GB")
        except (json.JSONDecodeError, KeyError):
            s.check("T7 ollama_setup --json", False, "输出不是合法 JSON")
            info = {}

        # T8: ollama 实体识别端到端（无 Ollama 则 SKIP）
        if info.get("ready"):
            draft = w / "t8.draft.json"
            r = run([PY, str(SKILL_DIR / "scripts" / "ollama_entities.py"), str(src),
                     "-o", str(draft), "--yes"])
            if r.returncode == 0 and draft.exists():
                originals = {e["original"] for e in json.loads(draft.read_text(encoding="utf-8"))["entities"]}
                hit = {"张测三", "李测四"} & originals
                blind = {"1985年3月12日出生", "浙A12345"} & originals
                s.check("T8 Ollama 识别（姓名+盲区字段）", len(hit) == 2 and len(blind) >= 1,
                        f"命中 {len(originals)} 个实体")
            else:
                s.check("T8 Ollama 识别", False, r.stderr[-200:])
        else:
            s.skip("T8 Ollama 识别", "本机未配置 Ollama，全本地路径不可用（在线路径不受影响）")

    report(s, args.json)
    sys.exit(0 if all(r["status"] != "FAIL" for r in s.results) else 1)


def report(s: Suite, as_json: bool) -> None:
    if as_json:
        print(json.dumps(s.results, ensure_ascii=False, indent=2))
        return
    for r in s.results:
        mark = {"PASS": "✅", "FAIL": "❌", "SKIP": "⏭️ "}[r["status"]]
        line = f"{mark} {r['name']}"
        if r["detail"]:
            line += f"  — {r['detail']}"
        print(line)
    n = {k: sum(1 for r in s.results if r["status"] == k) for k in ("PASS", "FAIL", "SKIP")}
    print(f"\n合计：{n['PASS']} PASS / {n['FAIL']} FAIL / {n['SKIP']} SKIP")


if __name__ == "__main__":
    main()
