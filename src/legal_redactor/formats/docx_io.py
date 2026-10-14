"""DOCX load / text extract / mapped rewrite via OOXML.

Handles track-changes (w:ins / w:del) and run-split entity strings by:
1. accepting insertions and dropping deletions;
2. applying replacements at run granularity when possible, else paragraph level.
Also cleans core/app/custom props, author/rsid attributes, and mailto/external rels.
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree

from ..entities import apply_mapping_to_text

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
XML_NS = "http://www.w3.org/XML/1998/namespace"
NSMAP = {"w": W_NS}
W = "{%s}" % W_NS

# Package parts that may carry visible body text
_TEXT_PART_RE = re.compile(
    r"^word/(document|header\d+|footer\d+|footnotes|endnotes|comments)\.xml$"
)
_RSID_ATTR_RE = re.compile(r"^rsid", re.IGNORECASE)


def _local(tag: str) -> str:
    if not isinstance(tag, str):
        return ""
    if tag.startswith("{"):
        return tag.rsplit("}", 1)[-1]
    return tag


REVISION_MODES = ("accept", "reject", "keep-and-redact")


def _unwrap(el: etree._Element) -> None:
    parent = el.getparent()
    if parent is None:
        return
    idx = parent.index(el)
    for offset, child in enumerate(list(el)):
        parent.insert(idx + offset, child)
    parent.remove(el)


def _flatten_revisions(root: etree._Element, revisions: str = "accept") -> None:
    """Resolve tracked changes.

    accept: keep insertions, drop deletions (default, historic behaviour).
    reject: drop insertions, restore deletions (w:delText -> w:t).
    keep-and-redact: leave w:ins / w:del in place; text inside both is redacted.
    """
    if revisions == "keep-and-redact":
        return
    if revisions not in REVISION_MODES:
        raise ValueError(f"revisions must be one of {REVISION_MODES}")
    drop = {"del", "moveFrom"} if revisions == "accept" else {"ins", "moveTo"}
    keep = {"ins", "moveTo"} if revisions == "accept" else {"del", "moveFrom"}
    for el in list(root.iter()):
        if _local(el.tag) in drop:
            parent = el.getparent()
            if parent is not None:
                parent.remove(el)
    if revisions == "reject":
        for el in root.iter():
            if _local(el.tag) == "delText":
                el.tag = f"{W}t"
            elif _local(el.tag) == "delInstrText":
                el.tag = f"{W}instrText"
    changed = True
    while changed:
        changed = False
        for el in list(root.iter()):
            if _local(el.tag) in keep and el.getparent() is not None:
                _unwrap(el)
                changed = True
                break


def _scrub_authors_and_rsids(root: etree._Element) -> None:
    for el in root.iter():
        # Attribute keys may be Clark notation
        for attr in list(el.attrib):
            local = _local(attr)
            if local in ("author", "lastModifiedBy"):
                el.attrib[attr] = "作者"
            elif local == "initials":
                el.attrib[attr] = "作"
            elif _RSID_ATTR_RE.match(local):
                del el.attrib[attr]


def _revision_owner(node: etree._Element) -> etree._Element | None:
    for anc in node.iterancestors():
        if _local(anc.tag) in ("ins", "del", "moveFrom", "moveTo"):
            return anc
        if _local(anc.tag) == "p":
            return None
    return None


def _paragraph_text(p: etree._Element) -> str:
    parts: list[str] = []
    owner: object = False
    for node in p.iter():
        if _local(node.tag) in ("t", "delText"):
            cur = _revision_owner(node)
            # keep-and-redact: deleted & inserted text must not glue together
            if owner is not False and cur is not owner and parts:
                parts.append("\n")
            owner = cur
            parts.append(node.text or "")
        elif _local(node.tag) == "tab":
            parts.append("\t")
        elif _local(node.tag) == "br":
            parts.append("\n")
    return "".join(parts)


def _preserve_space(t: etree._Element, text: str) -> None:
    if text[:1].isspace() or text[-1:].isspace():
        t.set(f"{{{XML_NS}}}space", "preserve")


def _set_paragraph_text(p: etree._Element, new_text: str) -> None:
    """Collapse paragraph visible text into the first w:t, keeping its run props."""
    t_nodes = [n for n in p.iter() if _local(n.tag) == "t"]
    if not t_nodes:
        r = etree.SubElement(p, f"{W}r")
        t = etree.SubElement(r, f"{W}t")
        t.text = new_text
        _preserve_space(t, new_text)
        return

    first = t_nodes[0]
    first.text = new_text
    _preserve_space(first, new_text)
    for t in t_nodes[1:]:
        t.text = ""


def _redact_paragraph(p: etree._Element, mapping: list[tuple[str, str]]) -> str:
    """Apply mapping; prefer per-run edits so mid-paragraph formatting survives."""
    original = _paragraph_text(p)
    if not original:
        return original
    updated = apply_mapping_to_text(original, mapping)
    if updated == original:
        return original

    t_nodes = [n for n in p.iter() if _local(n.tag) in ("t", "delText")]
    has_revision = any(_local(n.tag) in ("ins", "del") for n in p.iter())
    if t_nodes:
        for t in t_nodes:
            old = t.text or ""
            new = apply_mapping_to_text(old, mapping)
            if new != old:
                t.text = new
                _preserve_space(t, new)
        after = _paragraph_text(p)
        if after == updated:
            return updated
        if has_revision:
            # Collapsing would destroy tracked-change structure; report the
            # per-run result and let residual verify flag what is left.
            _REVISION_SPLIT_WARN.append(original[:40])
            return after
        # Entity spanned runs (or per-run order differed) — safe collapse.
        _set_paragraph_text(p, updated)
        return updated

    _set_paragraph_text(p, updated)
    return updated


_REVISION_SPLIT_WARN: list[str] = []


def _process_xml_bytes(
    data: bytes, mapping: list[tuple[str, str]] | None, revisions: str = "accept"
) -> tuple[bytes, str]:
    root = etree.fromstring(data)
    _flatten_revisions(root, revisions)
    if mapping is not None:
        _scrub_authors_and_rsids(root)

    texts: list[str] = []
    for p in root.iter():
        if _local(p.tag) != "p":
            continue
        original = _paragraph_text(p)
        if not original:
            continue
        if mapping is None:
            texts.append(original)
            continue
        texts.append(_redact_paragraph(p, mapping))

    out = etree.tostring(
        root,
        xml_declaration=True,
        encoding="UTF-8",
        standalone=True,
    )
    return out, "\n".join(texts)


def extract_text(path: Path, revisions: str = "accept") -> str:
    chunks: list[str] = []
    with zipfile.ZipFile(path, "r") as zin:
        names = [n for n in zin.namelist() if _TEXT_PART_RE.match(n)]
        names.sort(key=lambda n: (0 if n.endswith("document.xml") else 1, n))
        for name in names:
            _, text = _process_xml_bytes(zin.read(name), mapping=None, revisions=revisions)
            if text.strip():
                chunks.append(text)
    return "\n".join(chunks)


def _redact_core_props(data: bytes, mapping: list[tuple[str, str]]) -> bytes:
    try:
        root = etree.fromstring(data)
    except etree.XMLSyntaxError:
        return data
    for el in root.iter():
        local = _local(el.tag)
        if local in {"creator", "lastModifiedBy"}:
            el.text = "作者"
            continue
        if el.text and isinstance(el.text, str) and el.text.strip():
            el.text = apply_mapping_to_text(el.text, mapping)
        if el.tail and isinstance(el.tail, str) and el.tail.strip():
            el.tail = apply_mapping_to_text(el.tail, mapping)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _clean_app_props(data: bytes) -> bytes:
    try:
        root = etree.fromstring(data)
    except etree.XMLSyntaxError:
        return data
    for el in root.iter():
        local = _local(el.tag)
        if local in {"Company", "Manager", "CompanyDoc", "HyperlinkBase"}:
            el.text = ""
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _clean_custom_props(data: bytes) -> bytes:
    try:
        root = etree.fromstring(data)
    except etree.XMLSyntaxError:
        return data
    # Remove all property children under Properties
    for child in list(root):
        root.remove(child)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _scrub_misc_part(data: bytes) -> bytes:
    try:
        root = etree.fromstring(data)
    except etree.XMLSyntaxError:
        return data
    for el in list(root.iter()):
        if _local(el.tag) in {"rsids", "presenceInfo"}:
            parent = el.getparent()
            if parent is not None:
                parent.remove(el)
    _scrub_authors_and_rsids(root)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _scrub_rels(data: bytes) -> bytes:
    try:
        root = etree.fromstring(data)
    except etree.XMLSyntaxError:
        return data
    changed = False
    for el in root.iter():
        if _local(el.tag) != "Relationship":
            continue
        target = el.get("Target") or ""
        if target.lower().startswith("mailto:") or el.get("TargetMode") == "External":
            el.set("Target", "#")
            changed = True
    if not changed:
        return data
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def redact_docx(
    input_path: Path,
    output_path: Path,
    mapping: list[tuple[str, str]],
    warnings: list[str] | None = None,
    *,
    preserve_hyperlinks: bool = False,
    preserve_bookmarks: bool = False,  # unused for docx; accepted for API symmetry
    revisions: str = "accept",
    ocr_images: str = "auto",
    image_categories: set[str] | None = None,
) -> str:
    input_path = Path(input_path)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    collected: list[str] = []
    warn = warnings if warnings is not None else []
    with zipfile.ZipFile(input_path, "r") as zin:
        names = zin.namelist()
        if any(n.startswith("word/media/") for n in names):
            warn.extend(
                _media_warnings(input_path, mapping, ocr_images, image_categories)
            )
        if any(n.startswith("word/embeddings/") for n in names):
            warn.append("docx contains word/embeddings/* (OLE) — review or remove before sharing")
        with zipfile.ZipFile(output_path, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                name = item.filename
                if _TEXT_PART_RE.match(name):
                    data, text = _process_xml_bytes(data, mapping, revisions)
                    if text.strip():
                        collected.append(text)
                elif name == "docProps/core.xml":
                    data = _redact_core_props(data, mapping)
                elif name == "docProps/app.xml":
                    data = _clean_app_props(data)
                elif name == "docProps/custom.xml":
                    data = _clean_custom_props(data)
                elif name.startswith("word/_rels/") and name.endswith(".rels"):
                    if not preserve_hyperlinks:
                        data = _scrub_rels(data)
                elif name in ("word/settings.xml", "word/people.xml"):
                    data = _scrub_misc_part(data)
                zout.writestr(item, data)

    if _REVISION_SPLIT_WARN:
        warn.append(
            "keep-and-redact: entity spans a tracked-change boundary in "
            f"{len(_REVISION_SPLIT_WARN)} paragraph(s); not collapsed — check residual report"
        )
        _REVISION_SPLIT_WARN.clear()
    return "\n".join(collected)


def _media_warnings(
    path: Path,
    mapping: list[tuple[str, str]],
    ocr_images: str,
    categories: set[str] | None,
) -> list[str]:
    """List word/media images; OCR them when Tesseract is available (auto/on)."""
    from ..image_scan import ocr_available, scan_docx_media

    if ocr_images == "off" or not ocr_available():
        why = "--ocr-images off" if ocr_images == "off" else "Tesseract not installed"
        if ocr_images == "on" and why != "--ocr-images off":
            raise RuntimeError("--ocr-images on requires Tesseract (chi_sim) — see references/ocr.md")
        with zipfile.ZipFile(path) as z:
            media = [n for n in z.namelist() if n.startswith("word/media/")]
        return [
            f"docx contains {len(media)} embedded image(s) ({why}) — not OCR-scanned; "
            "review manually: " + ", ".join(media[:10])
        ]
    from ..patterns import categories_for_mode

    cats = categories if categories is not None else categories_for_mode("ai")
    items = scan_docx_media(path, categories=cats, extra_terms=[o for o, _ in mapping])
    out: list[str] = []
    for it in items:
        if it["hits"]:
            kinds = sorted({h["category"] for h in it["hits"]})
            out.append(
                f"IMAGE-PII {it['name']}: OCR found {len(it['hits'])} hit(s) ({', '.join(kinds)}) — "
                "image NOT modified; replace or crop it before sharing"
            )
        elif not it["ocr"]:
            out.append(f"docx image {it['name']} not OCR-scanned (undecodable/tiny/unsupported); review manually")
    if not out:
        out.append(f"docx word/media images OCR-scanned ({len(items)}), no structural PII found — still glance at seals/signatures")
    return out


def create_sample_docx(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for block in text.split("\n"):
        doc.add_paragraph(block)
    doc.save(str(path))


def create_mixed_run_docx(path: Path, *, plain_prefix: str, bold_secret: str, plain_suffix: str) -> None:
    """Fixture helper: one paragraph with a bold middle run (formatting preservation tests)."""
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    p = doc.add_paragraph()
    p.add_run(plain_prefix)
    bold = p.add_run(bold_secret)
    bold.bold = True
    p.add_run(plain_suffix)
    doc.save(str(path))
