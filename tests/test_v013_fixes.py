"""0.13.0: leaks found by testing realistic judgments, contracts and Word files."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from legal_redactor.cli import main
from legal_redactor.entities import apply_mapping_to_text, build_plan
from legal_redactor.ledger_store import restore_text
from legal_redactor.patterns import detect_structural
from legal_redactor.pipeline import redact_file
from legal_redactor.suspects import detect_suspects
from legal_redactor.verify import add_unreplaced, scan_residual, unreplaced_originals


def _texts(suspects, category=None):
    return [s.text for s in suspects if category is None or s.category == category]


# ---------- organizations ----------


def test_org_in_judgment_sentence_not_one_blob():
    text = (
        "原告星测影视文化传媒有限公司（以下简称星测公司）与被告杭州云测网络科技有限公司"
        "（以下简称云测公司）纠纷一案。"
    )
    orgs = _texts(detect_suspects(text), "organization")
    assert orgs == ["星测影视文化传媒有限公司", "杭州云测网络科技有限公司"]


def test_org_after_verb_clause_cut_at_role_word():
    text = "被告杭州云测网络科技有限公司于本判决生效之日起十日内赔偿原告星测影视文化传媒有限公司经济损失"
    assert _texts(detect_suspects(text), "organization") == [
        "杭州云测网络科技有限公司",
        "星测影视文化传媒有限公司",
    ]


def test_org_clause_prefix_dropped_when_clean_name_known():
    text = "甲方：星测文化传媒有限公司。\n本合同由星测文化传媒有限公司负责履行。"
    assert _texts(detect_suspects(text), "organization") == ["星测文化传媒有限公司"]


def test_unanchored_org_is_review_only():
    text = "这是一段没有任何标点和角色词的很长很长很长很长很长很长的叙述星测文化传媒有限公司"
    sus = [s for s in detect_suspects(text) if s.category == "organization"]
    assert sus and sus[0].reason == "org_suffix_unanchored"
    plan = build_plan(text, mode="ai", auto_confident=True)
    assert not [e for e in plan.entities if e.category == "organization"]


def test_unquoted_abbreviation_shares_alias():
    text = "原告：星测影视文化传媒有限公司。\n原告星测影视文化传媒有限公司（以下简称星测公司）诉称，星测公司享有权利。"
    plan = build_plan(text, mode="ai", auto_confident=True)
    by = {e.original: e.replacement for e in plan.entities}
    assert by["星测公司"] == by["星测影视文化传媒有限公司"]


# ---------- persons ----------


def test_common_words_after_role_are_not_names():
    text = "第三人陈述，被告明知，原告方面，原告经济损失100000元，证人陈述称，"
    assert _texts(detect_suspects(text), "person") == []


def test_no_colon_name_still_found():
    assert "张测三" in _texts(detect_suspects("原告张测三，男，1985年出生"), "person")


def test_narrative_relation_names():
    text = "证人王测娟出庭作证称，王测娟的丈夫刘测胜亦知情。涉案APP由其表弟周测川实际运营。"
    names = _texts(detect_suspects(text), "person")
    assert {"王测娟", "刘测胜", "周测川"} <= set(names)


def test_narrative_names_kept_in_production_except_witness():
    text = "证人王测娟出庭作证。其表弟周测川实际运营。"
    plan = build_plan(text, mode="production", auto_confident=True)
    replaced = {e.original for e in plan.entities}
    assert "周测川" not in replaced  # kin = role other → production keeps


def test_court_alias_numbered_for_same_surname():
    text = "原告：张测三，男。\n被告：张测四，男。\n张测三与张测四系兄弟。"
    plan = build_plan(text, mode="ai", auto_confident=True, alias_style="court")
    by = {e.original: e.replacement for e in plan.entities if e.category == "person"}
    assert by == {"张测三": "张某1", "张测四": "张某2"}


# ---------- accounts ----------


def test_account_label_without_colon_and_in_table():
    for text in ("开户行为测试银行，账号0200004509200123456。", "收款账户\n0200004509200123456"):
        hits = [h for h in detect_structural(text) if h.category == "bank_account"]
        assert hits and hits[0].auto_replace, text


# ---------- verification ----------


def test_ai_verify_lists_names_left_in_output():
    out = "原告：星测影视文化传媒有限公司。法定代表人：某甲。"
    r = scan_residual(out, mode="ai", nl=True)
    assert r.status == "WARN"
    assert any(w["text"] == "星测影视文化传媒有限公司" for w in r.warnings)
    assert scan_residual(out, mode="production", nl=True).status == "PASS"


def test_court_staff_not_flagged_by_default():
    r = scan_residual("审判长 赵测审\n书记员 钱测记", mode="ai", nl=True)
    assert r.status == "PASS"


def test_unreplaced_original_is_fail():
    mapping = [("陈测国", "某甲")]
    r = scan_residual("被告某甲，陈测国", mode="ai")
    add_unreplaced(r, unreplaced_originals("被告某甲，陈测国", mapping))
    assert r.status == "FAIL" and r.hits[-1]["category"] == "unreplaced_entity"
    # a wrapped PDF line still counts
    assert unreplaced_originals("陈测\n国", mapping, join_lines=True) == ["陈测国"]


def test_cli_verify_strict_exit_codes(tmp_path: Path):
    leaked = tmp_path / "leak.txt"
    leaked.write_text("原告：星测影视文化传媒有限公司。", encoding="utf-8")
    assert main(["verify", str(leaked), "--mode", "ai"]) == 0
    assert main(["verify", str(leaked), "--mode", "ai", "--strict"]) == 1
    allow = tmp_path / "allow.txt"
    allow.write_text("星测影视文化传媒有限公司\n", encoding="utf-8")
    assert main(["verify", str(leaked), "--mode", "ai", "--strict", "--allowlist", str(allow)]) == 0


# ---------- formats ----------


def test_docx_split_runs_keep_formatting(tmp_path: Path):
    src = tmp_path / "runs.docx"
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("被告：").bold = True
    p.add_run("陈")
    p.add_run("测国，男，联系电话")
    p.add_run("138")
    p.add_run("12345678。")
    doc.save(str(src))
    res = redact_file(src, mode="ai", output_path=tmp_path / "out.docx", auto_confident=True)
    runs = Document(str(res.output_path)).paragraphs[0].runs
    assert runs[0].text == "被告：" and runs[0].bold
    assert not any(r.bold for r in runs[1:])
    text = "".join(r.text for r in runs)
    assert "陈" not in text and "13812345678" not in text and "[手机号]" in text


def test_text_output_keeps_line_endings(tmp_path: Path):
    for eol in ("\n", "\r\n"):
        src = tmp_path / f"in{len(eol)}.txt"
        src.write_bytes(f"原告：张测三{eol}手机13812345678{eol}".encode("utf-8"))
        res = redact_file(src, mode="ai", output_path=tmp_path / f"out{len(eol)}.txt")
        data = res.output_path.read_bytes()
        assert data.count(eol.encode()) == 2
        if eol == "\n":
            assert b"\r\n" not in data


def test_mapping_spans_match_longest_first():
    mapping = [("星测文化传媒有限公司", "某单位A"), ("星测公司", "某单位A"), ("星测", "X")]
    assert apply_mapping_to_text("星测文化传媒有限公司即星测公司，星测", mapping) == "某单位A即某单位A，X"


# ---------- restore ----------


def test_restore_refuses_bucket_placeholders():
    ledger = {"entities": [
        {"original": "100000元", "replacement": "数十万元", "category": "amount"},
        {"original": "300000元", "replacement": "数十万元", "category": "amount"},
        {"original": "陈测国", "replacement": "某甲", "category": "person"},
    ]}
    out, warns = restore_text("某甲应赔偿数十万元", ledger)
    assert out == "陈测国应赔偿数十万元"
    assert any("NOT restored" in w for w in warns)
