"""Tests for the local-first web server and endpoints."""

from __future__ import annotations

import io
import json
import socket
import tempfile
import threading
from http.server import ThreadingHTTPServer
from pathlib import Path
from urllib.request import Request, urlopen

import docx
import pymupdf as fitz
import pytest

from legal_redactor.web import (
    WebRequestHandler,
    _find_available_port,
    _safe_filename,
    parse_multipart,
)


def _build_multipart(fields: dict[str, str], files: dict[str, tuple[str, bytes, str]]) -> tuple[str, bytes]:
    boundary = "----WebKitFormBoundaryTest12345"
    lines: list[bytes] = []
    for k, v in fields.items():
        lines.append(f"--{boundary}\r\n".encode("utf-8"))
        lines.append(f'Content-Disposition: form-data; name="{k}"\r\n\r\n'.encode("utf-8"))
        lines.append(f"{v}\r\n".encode("utf-8"))

    for k, (fname, content, ctype) in files.items():
        lines.append(f"--{boundary}\r\n".encode("utf-8"))
        lines.append(
            f'Content-Disposition: form-data; name="{k}"; filename="{fname}"\r\n'.encode("utf-8")
        )
        lines.append(f"Content-Type: {ctype}\r\n\r\n".encode("utf-8"))
        lines.append(content)
        lines.append(b"\r\n")

    lines.append(f"--{boundary}--\r\n".encode("utf-8"))
    return f"multipart/form-data; boundary={boundary}", b"".join(lines)


@pytest.fixture(scope="module")
def web_server():
    port = _find_available_port(8890, max_attempts=20)
    server = ThreadingHTTPServer(("127.0.0.1", port), WebRequestHandler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    base_url = f"http://127.0.0.1:{port}"
    yield base_url
    server.shutdown()
    server.server_close()


def test_safe_filename():
    assert _safe_filename("test.pdf") == "test.pdf"
    assert _safe_filename("../../bad/path/file.docx") == "file.docx"
    assert "民事起诉状.pdf" in _safe_filename("民事起诉状.pdf")


def test_parse_multipart():
    ctype, body = _build_multipart(
        {"mode": "ai"},
        {"file": ("民事起诉状.pdf", b"%PDF-1.4 mock", "application/pdf")},
    )
    form, files = parse_multipart(ctype, body)
    assert form["mode"] == "ai"
    assert files["file"]["filename"] == "民事起诉状.pdf"
    assert files["file"]["content"] == b"%PDF-1.4 mock"


def test_get_root(web_server):
    req = Request(f"{web_server}/")
    with urlopen(req) as resp:
        assert resp.status == 200
        html = resp.read().decode("utf-8")
        assert "中国法律文书" in html
        assert "文档脱敏" in html


def test_get_doctor(web_server):
    req = Request(f"{web_server}/api/doctor")
    with urlopen(req) as resp:
        assert resp.status == 200
        data = json.loads(resp.read().decode("utf-8"))
        assert "version" in data
        assert "tesseract" in data
        assert "opencv" in data


def test_redact_text_and_restore(web_server):
    sample_text = "原告张三，身份证号110101199003072345，电话13800138000。"
    payload = json.dumps({
        "text": sample_text,
        "mode": "ai",
        "auto_confident": True,
    }).encode("utf-8")

    req = Request(
        f"{web_server}/api/redact-text",
        data=payload,
        headers={"Content-Type": "application/json"},
    )
    with urlopen(req) as resp:
        assert resp.status == 200
        data = json.loads(resp.read().decode("utf-8"))
        assert data["ok"] is True
        redacted = data["redacted_text"]
        assert "某甲" in redacted
        assert "[身份证号]" in redacted
        assert "[手机号]" in redacted
        ledger = data["ledger"]

    # Now restore
    ai_reply = f"经查，{redacted}主体适格。"
    restore_payload = json.dumps({
        "text": ai_reply,
        "ledger": ledger,
    }).encode("utf-8")
    req_restore = Request(
        f"{web_server}/api/restore-text",
        data=restore_payload,
        headers={"Content-Type": "application/json"},
    )
    with urlopen(req_restore) as resp:
        assert resp.status == 200
        restore_data = json.loads(resp.read().decode("utf-8"))
        assert restore_data["ok"] is True
        assert "张三" in restore_data["restored_text"]


def test_inspect_and_redact_pdf(web_server):
    # 1. Create a test PDF with text layer
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "原告李四，男，身份证号330102199001011234，电话13900139000。")
    pdf_bytes = doc.tobytes()
    doc.close()

    # 2. Inspect PDF
    ctype, body = _build_multipart({}, {"file": ("起诉状.pdf", pdf_bytes, "application/pdf")})
    req_inspect = Request(f"{web_server}/api/inspect", data=body, headers={"Content-Type": ctype})
    with urlopen(req_inspect) as resp:
        assert resp.status == 200
        inspect_data = json.loads(resp.read().decode("utf-8"))
        assert inspect_data["ok"] is True
        assert inspect_data["is_pdf"] is True
        assert inspect_data["has_text_layer"] is True
        assert inspect_data["page_count"] == 1

    # 3. Redact PDF
    ctype_redact, body_redact = _build_multipart(
        {"mode": "ai", "auto_confident": "true"},
        {"file": ("起诉状.pdf", pdf_bytes, "application/pdf")},
    )
    req_redact = Request(f"{web_server}/api/redact-file", data=body_redact, headers={"Content-Type": ctype_redact})
    with urlopen(req_redact) as resp:
        assert resp.status == 200
        redact_data = json.loads(resp.read().decode("utf-8"))
        assert redact_data["ok"] is True
        assert redact_data["status"] == "PASS"
        assert "download_file_url" in redact_data

        dl_file_url = f"{web_server}{redact_data['download_file_url']}"
        dl_audit_url = f"{web_server}{redact_data['download_audit_url']}"
        dl_ledger_url = f"{web_server}{redact_data['download_ledger_url']}"

    # 4. Download file and verify
    with urlopen(Request(dl_file_url)) as resp:
        assert resp.status == 200
        out_pdf_bytes = resp.read()
        out_doc = fitz.open(stream=out_pdf_bytes, filetype="pdf")
        text = out_doc[0].get_text("text")
        out_doc.close()
        assert "李四" not in text

    # 5. Download audit
    with urlopen(Request(dl_audit_url)) as resp:
        assert resp.status == 200
        audit_content = resp.read().decode("utf-8")
        assert "Audit Report" in audit_content or "审计" in audit_content or "PASS" in audit_content

    # 6. Download ledger
    with urlopen(Request(dl_ledger_url)) as resp:
        assert resp.status == 200
        ledger_data = json.loads(resp.read().decode("utf-8"))
        assert "entities" in ledger_data or "plan" in ledger_data


def test_redact_docx(web_server):
    # 1. Create a test docx
    doc = docx.Document()
    doc.add_paragraph("原告王五，统一社会信用代码91330100MA2B123456。")
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    # 2. Redact docx with keep_categories=uscc
    ctype, body = _build_multipart(
        {"mode": "production", "keep_categories": "uscc", "auto_confident": "true"},
        {"file": ("合同.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    req = Request(f"{web_server}/api/redact-file", data=body, headers={"Content-Type": ctype})
    with urlopen(req) as resp:
        assert resp.status == 200
        data = json.loads(resp.read().decode("utf-8"))
        assert data["ok"] is True
        dl_url = f"{web_server}{data['download_file_url']}"

    # 3. Download docx
    with urlopen(Request(dl_url)) as resp:
        assert resp.status == 200
        out_docx_bytes = resp.read()
        out_doc = docx.Document(io.BytesIO(out_docx_bytes))
        paragraphs = [p.text for p in out_doc.paragraphs]
        full_text = "\n".join(paragraphs)
        # Production mode preserves party王五 and keep_categories preserves uscc
        assert "王五" in full_text
        assert "91330100MA2B123456" in full_text
